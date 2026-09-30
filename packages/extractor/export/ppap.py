"""
PPAP dimensional results workbook (AIAG style): item, specification, limits, result, OK / NOT OK.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage

from .as9102 import BORDER, CELL_FONT, CENTER, FAIL_FONT, PASS_FONT, TITLE_FONT, WRAP, _header_row, _labelled
from .characteristics import balloon_label, derive_limits, display_status, fmt, num, requirement_text, sort_key
from .fcf_image import render_fcf_png


def build_ppap_workbook(info: dict[str, Any], chars: list[dict[str, Any]], defaults: dict[str, float]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dimensional Results"
    labels = ["Item", "Dimension / Specification", "Min", "Max", "Result", "OK / Not OK", "Zone", "Sheet", "Comments"]
    widths = [8, 36, 12, 12, 12, 14, 8, 8, 30]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(labels))
    ws.cell(row=1, column=1, value="PPAP Dimensional Results").font = TITLE_FONT
    _labelled(ws, 2, 1, "Part Number", info.get("partNumber", ""), 2)
    _labelled(ws, 2, 3, "Part Name", info.get("partName", ""), 2)
    _labelled(ws, 2, 5, "Supplier", info.get("organization", ""), 2)
    _labelled(ws, 2, 7, "Drawing Rev", info.get("drawingRevision", ""), 1)
    _labelled(ws, 2, 8, "Date", date.today().isoformat(), 2)
    _header_row(ws, 3, labels, widths)
    row = 4
    for c in sorted(chars, key=sort_key):
        limits = derive_limits(c, defaults)
        status = display_status(c, limits)
        result = c.get("result")
        n = num(result)
        p = int(c.get("places") or 0)
        result_txt = fmt(n, p) if (n is not None and not isinstance(result, str)) else ("" if result is None else str(result))
        verdict = "OK" if status == "Pass" else "NOT OK" if status == "Fail" else ""
        values = [
            balloon_label(c), requirement_text(c, limits), fmt(limits.min, p), fmt(limits.max, p), result_txt, verdict,
            c.get("zone") or "", int(c.get("page") or 0) + 1, c.get("comments") or "",
        ]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = PASS_FONT if (i == 6 and verdict == "OK") else FAIL_FONT if (i == 6 and verdict == "NOT OK") else CELL_FONT
            cell.border = BORDER
            cell.alignment = CENTER if i in (1, 3, 4, 5, 6, 7, 8) else WRAP
        if c.get("toleranceType") == "GD&T" and c.get("gdt"):
            img = XlImage(io.BytesIO(render_fcf_png(c["gdt"], height_px=22)))
            img.width, img.height = img.width // 2, 22
            ws.cell(row=row, column=2).value = None  # the frame image replaces the text
            ws.add_image(img, f"B{row}")
            ws.row_dimensions[row].height = 24
        row += 1
    ws.freeze_panes = "A4"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
