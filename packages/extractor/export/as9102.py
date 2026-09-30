"""
AS9102 Rev C workbook: Form 1 (Part Number Accountability), Form 2 (Product Accountability),
Form 3 (Characteristic Accountability). GD&T frames are placed in cell 8 as images.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .characteristics import balloon_label, derive_limits, display_status, fmt, num, requirement_text, sort_key
from .fcf_image import render_fcf_png

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(name="Arial", size=12, bold=True)
LABEL_FONT = Font(name="Arial", size=8, bold=True)
CELL_FONT = Font(name="Arial", size=10)
PASS_FONT = Font(name="Arial", size=10, bold=True, color="0A7A3C")
FAIL_FONT = Font(name="Arial", size=10, bold=True, color="B3261E")
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _labelled(ws, row: int, col: int, label: str, value: str, width_cols: int = 1) -> None:
    """Field cell in the AS9102 style: small bold label on top, value below, one bordered box."""
    c = ws.cell(row=row, column=col, value=f"{label}\n{value}")
    c.font = CELL_FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.border = BORDER
    if width_cols > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + width_cols - 1)
        for k in range(1, width_cols):
            ws.cell(row=row, column=col + k).border = BORDER
    ws.row_dimensions[row].height = 30


def _header_row(ws, row: int, labels: list[str], widths: list[int]) -> None:
    for i, (label, width) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = LABEL_FONT
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = CENTER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 32


def _part_header(ws, info: dict[str, Any], title: str, cols: int) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.row_dimensions[1].height = 22
    _labelled(ws, 2, 1, "1. Part Number", info.get("partNumber", ""), 3)
    _labelled(ws, 2, 4, "2. Part Name", info.get("partName", ""), 3)
    _labelled(ws, 2, 7, "3. Serial Number", info.get("serialNumbers", ""), 2)
    _labelled(ws, 2, 9, "4. FAIR Identifier", info.get("fairNumber", ""), max(1, cols - 8))
    return 3


def build_form1(ws, info: dict[str, Any], stats: dict[str, Any]) -> None:
    ws.title = "Form 1"
    ws.merge_cells("A1:F1")
    ws.cell(row=1, column=1, value="AS9102 Rev C  Form 1: Part Number Accountability").font = TITLE_FONT
    for col in "ABCDEF":
        ws.column_dimensions[col].width = 26
    fields = [
        ("1. Part Number", info.get("partNumber", "")),
        ("2. Part Name", info.get("partName", "")),
        ("3. Serial Number", info.get("serialNumbers", "")),
        ("4. FAIR Identifier", info.get("fairNumber", "")),
        ("5. Part Revision Level", info.get("partRevision", "")),
        ("6. Drawing Number", info.get("drawingNumber", "")),
        ("7. Drawing Revision Level", info.get("drawingRevision", "")),
        ("8. Additional Changes", info.get("additionalChanges", "")),
        ("9. Manufacturing Process Reference", info.get("processRef", "")),
        ("10. Organization Name", info.get("organization", "")),
        ("11. Supplier Code", info.get("supplierCode", "")),
        ("12. P.O. Number", info.get("poNumber", "")),
        ("13. Detail / Assembly FAI", info.get("detailOrAssembly", "Detail")),
        ("14. Full / Partial FAI", info.get("fullOrPartial", "Full")),
        ("15. Baseline Part Number", info.get("baselinePartNumber", "")),
        ("Customer", info.get("customer", "")),
        ("Characteristics on Form 3", str(stats.get("count", 0))),
        ("Results recorded", str(stats.get("measured", 0))),
        ("Nonconforming characteristics", str(stats.get("fails", 0))),
        ("Report generated", date.today().isoformat()),
    ]
    row = 2
    for i in range(0, len(fields), 3):
        for j, (label, value) in enumerate(fields[i : i + 3]):
            _labelled(ws, row, 1 + j * 2, label, value, 2)
        row += 1
    row += 1
    for j, label in enumerate(["19. Signature", "20. Date", "21. Reviewed By", "22. Date", "23. Customer Approval", "24. Date"]):
        _labelled(ws, row, 1 + j, label, "", 1)
        ws.column_dimensions[get_column_letter(1 + j)].width = 26


def build_form2(ws, info: dict[str, Any], rows: list[dict[str, Any]]) -> None:
    ws.title = "Form 2"
    labels = [
        "5. Material or Process Name", "6. Specification Number", "7. Code", "8. Special Process Supplier Code",
        "9. Customer Approval Verification", "10. Certificate of Conformance Number", "11. Functional Test Procedure Number",
        "12. Acceptance Report Number", "13. Comments", "14. Prepared By", "15. Date",
    ]
    widths = [28, 22, 8, 16, 14, 20, 18, 18, 24, 16, 12]
    row = _part_header(ws, info, "AS9102 Rev C  Form 2: Product Accountability", len(labels))
    _header_row(ws, row, labels, widths)
    row += 1
    code_for = {"Material": "M", "Special process": "SP", "Test": "T"}
    for r in rows:
        values = [
            r.get("name", ""), r.get("specNumber", ""), r.get("code") or code_for.get(r.get("type", ""), ""), r.get("supplierCode", ""),
            r.get("customerApproval", ""), r.get("cocNumber", ""), r.get("testProcedure", ""), r.get("acceptanceReport", ""),
            r.get("comments", ""), r.get("preparedBy", ""), r.get("date", ""),
        ]
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.font = CELL_FONT
            c.border = BORDER
            c.alignment = WRAP
        row += 1
    for _ in range(max(0, 3 - len(rows))):
        for i in range(1, len(labels) + 1):
            ws.cell(row=row, column=i).border = BORDER
        row += 1


def build_form3(ws, info: dict[str, Any], chars: list[dict[str, Any]], defaults: dict[str, float], sheet_names: dict[int, str] | None = None) -> dict[str, Any]:
    ws.title = "Form 3"
    labels = [
        "5. Char No.", "6. Reference Location", "7. Characteristic Designator", "8. Requirement", "8a. UoM",
        "9. Results", "10. Designed / Qualified Tooling", "11. Nonconformance Number", "14. Additional Data / Comments", "Status",
    ]
    widths = [10, 20, 16, 34, 8, 14, 18, 18, 30, 10]
    row = _part_header(ws, info, "AS9102 Rev C  Form 3: Characteristic Accountability, Verification and Compatibility Evaluation", len(labels))
    _header_row(ws, row, labels, widths)
    row += 1
    measured = fails = 0
    for c in sorted(chars, key=sort_key):
        limits = derive_limits(c, defaults)
        status = display_status(c, limits)
        result = c.get("result")
        if result not in (None, ""):
            measured += 1
        if status == "Fail":
            fails += 1
        zone = (c.get("zone") or "").strip()
        page = int(c.get("page") or 0) + 1
        ref = f"Sheet {page} Zone {zone}" if zone else f"Sheet {page}"
        n = num(result)
        result_txt = fmt(n, int(c.get("places") or 0)) if (n is not None and not isinstance(result, str)) else ("" if result is None else str(result))
        values = [
            balloon_label(c), ref, c.get("designator") or "", requirement_text(c, limits), c.get("units", ""),
            result_txt, c.get("tooling") or "", c.get("nonconformance") or "", c.get("comments") or "", status if status in ("Pass", "Fail") else "",
        ]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = PASS_FONT if (i in (6, 10) and status == "Pass") else FAIL_FONT if (i in (6, 10) and status == "Fail") else CELL_FONT
            cell.border = BORDER
            cell.alignment = CENTER if i in (1, 5, 6, 10) else WRAP
        if c.get("toleranceType") == "GD&T" and c.get("gdt"):
            # Rendered at 2x for crispness, shown at 22 px tall inside the row.
            img = XlImage(io.BytesIO(render_fcf_png(c["gdt"], height_px=22, scale=2)))
            img.width, img.height = img.width // 2, 22
            ws.cell(row=row, column=4).value = None  # the frame image replaces the text
            ws.add_image(img, f"D{row}")
            ws.row_dimensions[row].height = 24
        else:
            ws.row_dimensions[row].height = 18
        row += 1
    ws.freeze_panes = ws.cell(row=4, column=1)
    return {"count": len(chars), "measured": measured, "fails": fails}


def build_as9102_workbook(info: dict[str, Any], chars: list[dict[str, Any]], product_rows: list[dict[str, Any]], defaults: dict[str, float]) -> bytes:
    wb = Workbook()
    ws1 = wb.active
    ws2 = wb.create_sheet()
    ws3 = wb.create_sheet()
    stats = build_form3(ws3, info, chars, defaults)
    build_form2(ws2, info, product_rows)
    build_form1(ws1, info, stats)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
